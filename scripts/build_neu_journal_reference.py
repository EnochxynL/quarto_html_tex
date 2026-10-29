"""Prepare Pandoc style aliases from the NEU source without rebuilding its body.

First convert the legacy DOC read-only with Word to the build directory.
Only reference-xml.docx is published under styles/word.
"""
from copy import deepcopy
import json
import subprocess
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from build_word_references import ROOT, flatten_style, lookup, replace_child
from build_word_xml_references import build

WORK = ROOT/'build/word-xml-review/neu-journal'

def promote(source, seed, name, paragraph):
    try:
        style = lookup(seed,name)
    except StopIteration:
        style = seed.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    for kind in ['pPr','rPr']:
        props = flatten_style(source,paragraph.style,kind)
        direct = paragraph._p.find(qn('w:pPr')) if kind=='pPr' else None
        if kind=='rPr':
            run=max((r for r in paragraph.runs if r.text.strip()),key=lambda r:len(r.text),default=None)
            direct = run._r.find(qn('w:rPr')) if run is not None else None
        if direct is not None:
            for child in direct:
                if child.tag not in [qn('w:pStyle'),qn('w:sectPr'),qn('w:rPr')]:
                    replace_child(props,child)
        for tag in ['numPr','sectPr','pStyle','pageBreakBefore','rPr']:
            for child in list(props.findall(qn('w:'+tag))):
                props.remove(child)
        replace_child(style.element,props)
    return style

def main():
    source=Document(WORK/'source-converted.docx')
    default=WORK/'pandoc-default.docx'
    default.write_bytes(subprocess.check_output(['quarto','pandoc','--print-default-data-file','reference.docx']))
    seed=Document(default)
    # Use source defaults as well as inherited and direct paragraph/run formatting.
    defaults=source.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        replace_child(seed.styles.element,defaults)
    roles={'Title':0,'Author':1,'Abstract':3,'English Title':4,'English Author':5,
        'English Abstract':6,'Normal':9,'Body Text':9,'First Paragraph':9,'Compact':9,
        'Heading 1':11,'Heading 2':28,'Heading 3':28,
        'Caption':24,'Image Caption':24,'Table Caption':24,
        'Bibliography':101}
    for name,index in roles.items():
        promote(source,seed,name,source.paragraphs[index])
    lookup(seed,'Table Caption').paragraph_format.keep_with_next=True
    compact=lookup(seed,'Compact')
    compact.paragraph_format.first_line_indent=Pt(0)
    compact.paragraph_format.left_indent=Pt(0)
    compact.paragraph_format.right_indent=Pt(0)
    compact.element.find(qn('w:pPr')).find(qn('w:ind')).set(qn('w:firstLineChars'),'0')
    table=lookup(seed,'Table')
    table_props=table.element.find(qn('w:tblPr'))
    if table_props is None:
        table_props=OxmlElement('w:tblPr');table.element.append(table_props)
    width=OxmlElement('w:tblW');width.set(qn('w:type'),'pct');width.set(qn('w:w'),'5000')
    replace_child(table_props,width)
    # Source has no code/table examples: retain Pandoc baseline for these roles.
    try:
        code=lookup(seed,'Source Code')
    except StopIteration:
        code=seed.styles.add_style('Source Code',WD_STYLE_TYPE.PARAGRAPH)
    code.font.name='Consolas';code.font.size=Pt(9)
    code.paragraph_format.first_line_indent=Pt(0)
    for name in ['Figure','Captioned Figure']:
        lookup(seed,name).paragraph_format.first_line_indent=Pt(0)
    seed_path=WORK/'style-seed.docx'
    seed.save(seed_path)
    specimen=WORK/'specimen.md'
    specimen.write_text('''# 研究目标与方法

这是一份参考文档兼容性测试。正文、标题、表格、列表、公式及代码使用Pandoc语义样式。

## 方法说明

保留原始文档的内容、分节和单双栏设置，并新增输出样式。

| 项目 | 数值 |
|:-----|-----:|
| 测试A | 0.125 |
| 测试B | 0.250 |

: 测试表格

1. 数据准备。
2. 分析与验证。

$$ L=\\frac{1}{N}\\sum_{i=1}^{N}(y_i-\\hat y_i)^2 $$

```python
def mean(values):
    return sum(values) / len(values)
```

参考链接：[Quarto](https://quarto.org)。
''',encoding='utf-8')
    result=build('neu-journal',WORK/'source-converted.docx',seed=seed_path,specimen=specimen)
    result['legacy_source']='D:/MY_DATA/Documents/WeChat Files/wxid_vcrkfp9afe6212/FileStorage/File/2024-04/大作业模板(1).doc'
    (WORK/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2))

if __name__=='__main__':
    main()
