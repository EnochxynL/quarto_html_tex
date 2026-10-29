"""Deterministic document assembly above Pandoc's reference-doc style layer.

No fixed experiment data or heading-text matches: metadata supplies front matter,
and the optional landscape region is identified by semantic marker paragraphs.
"""
from pathlib import Path
from copy import deepcopy
import json,re
from lxml import etree
from docx import Document
from docx.oxml import OxmlElement,parse_xml
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm,Pt
from build_word_references import lookup,replace_child

def xml(node):return etree.tostring(node,encoding='unicode')
def section_xml(section):
 node=deepcopy(section._sectPr)
 for child in list(node):
  if child.tag in [qn('w:headerReference'),qn('w:footerReference'),qn('w:pgNumType'),qn('w:titlePg')]:node.remove(child)
 return xml(node)

def extract_layout(key,source,destination):
 d=Document(source)
 table_index={'python-experiment':0,'cad-report':1,'course-paper':3}.get(key)
 table=d.tables[table_index] if table_index is not None else None
 config=dict(version=2,name=key,
  front_matter={'python-experiment':'experiment-cover','cad-report':'group-cover','course-paper':'paper-cover','electronic-senses':'bilingual-inline'}[key],
  sections=[section_xml(s) for s in d.sections],
  main_columns=2 if key=='electronic-senses' else 1,
  chapter_new_page=key=='python-experiment',
  landscape_regions=key=='course-paper',
  table_properties=xml(table._tbl.tblPr) if table is not None else None,
  source_running_header=bool(''.join(p.text for s in d.sections for p in s.header.paragraphs).strip()),
  source_running_footer=bool(any(s.footer._element.xpath('.//w:t | .//w:fldSimple | .//w:instrText') for s in d.sections)),
  original_section_count=len(d.sections),
  notes=['reference-doc supplies styles; this layout adapter supplies structure.',
         'Original personal data and experiment text are replaced by specimen metadata.'])
 destination.write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf8')
 return config

def clean_section(text,kind=None,columns=None,restart=None):
 node=parse_xml(text.encode())
 if kind:
  n=OxmlElement('w:type');n.set(qn('w:val'),kind);replace_child(node,n)
 if columns:
  c=node.find(qn('w:cols'))
  if c is None:c=OxmlElement('w:cols');node.append(c)
  c.set(qn('w:num'),str(columns))
 if restart is not None:
  n=OxmlElement('w:pgNumType');n.set(qn('w:start'),str(restart));node.append(n)
 return node

def add_boundary_before(element,section):
 p=OxmlElement('w:p');props=OxmlElement('w:pPr')
 spacing=OxmlElement('w:spacing');spacing.set(qn('w:before'),'0');spacing.set(qn('w:after'),'0');spacing.set(qn('w:line'),'20');spacing.set(qn('w:lineRule'),'exact');props.append(spacing)
 props.append(deepcopy(section));p.append(props);element.addprevious(p);return p

def apply(input_docx,output_docx,layout,metadata):
 doc=Document(input_docx);body=doc.element.body;cfg=json.loads(Path(layout).read_text(encoding='utf8'));meta=json.loads(Path(metadata).read_text(encoding='utf8'))
 # Drop Pandoc's stock front matter; the source-derived scheme will assemble it.
 original=list(body);first=next(i+1 for i,n in enumerate(original) if n.tag==qn('w:p') and n.xpath('./w:pPr/w:pStyle[@w:val="LayoutBodyStart"]'))
 for node in original[:first]:body.remove(node)
 body_start=body[0]
 for sec in list(body.findall(qn('w:sectPr'))):body.remove(sec)
 sections=cfg['sections'];base=clean_section(sections[0],columns=1)
 body.append(deepcopy(base))
 prefix=[]
 def para(text='',style='Normal',center=False,bold=None):
  p=doc.add_paragraph(style=lookup(doc,style));p.add_run(text)
  if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
  if bold is not None:
   for run in p.runs:run.bold=bold
  prefix.append(p._p);return p
 def page():
  p=para();p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0);p.add_run().add_break()
  # A page break is deliberately a layout operation, not shared-body content.
  p.runs[-1]._r[-1].set(qn('w:type'),'page')
 def member_table():
  table=doc.add_table(rows=1,cols=4);prefix.append(table._tbl)
  for cell,text in zip(table.rows[0].cells,['姓名','学号','承担工作','工作比例']):cell.text=text
  for member in meta['members']:
   for cell,text in zip(table.add_row().cells,member):cell.text=text
  return table
 def toc():
  page();p=para('目录','TOC Heading',True)
  field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u');para()._p.append(field)
 def abstract(standalone=False):
  if standalone:para('摘要','Heading 1')
  para(meta['abstract'],'Body Text')
  para('关键词：'+meta['keywords'],'Body Text')
 front=cfg['front_matter']
 if front=='bilingual-inline':
  para(meta['title'],'Title')
  para(meta['author'],'Author')
  p=para(style='Abstract');p.add_run('摘  要：').bold=True;p.add_run(meta['abstract'])
  para(meta['title-en'],'EnglishTitle')
  para(meta['author-en'],'EnglishAuthor')
  p=para(style='EnglishAbstract');p.add_run('Abstract: ').bold=True;p.add_run(meta['abstract-en'])
 elif front=='experiment-cover':
  para('\n\n');para(meta['course'],'CoverCourse',True);para(meta['report-kind'],'CoverKind',True)
  para('\n');para(meta['title'],'CoverTopic',True);para('\n\n')
  for text in [meta['group'],'学号：00000001','姓名：'+meta['author'],'报告日期：'+meta['date']]:para(text,'CoverGroup',True)
  para('\n\n');para('本样张用于核对模板的版面与样式，示例内容不代表实际实验。','CoverNote')
  page()
 else:
  if front=='paper-cover':para(meta['institution'],'CoverInstitution',True)
  para(meta['course'],'CoverCourse',True)
  para(meta['paper-kind'] if front=='paper-cover' else meta['report-kind'],'CoverKind',True)
  para(meta['title'],'CoverTopic',True);para(meta['group'],'CoverGroup',True)
  member_table()
  if front=='paper-cover':para('指导教师：'+meta['supervisor'],'CoverSupervisor',True)
  para(meta['date'],'CoverDate',True)
  toc()
  if front=='paper-cover':
   # Separate unnumbered cover/TOC from numbered abstract and main matter.
   for n in prefix:body.remove(n)
   for n in prefix:body_start.addprevious(n)
   prefix=[]
   add_boundary_before(body_start,clean_section(sections[0],columns=1))
   abstract(True);page()
  else:page()
 for n in prefix:body.remove(n)
 for n in prefix:body_start.addprevious(n)
 # Section transitions: source compact bilingual header -> double-column body.
 if cfg['main_columns']==2:
  add_boundary_before(body_start,clean_section(sections[0],columns=1))
  current=clean_section(sections[1],kind='continuous',columns=2)
  # Source returns to one column for the ending bibliography region.
  bib=next((p._p for p in doc.paragraphs if p._p.xpath('./w:pPr/w:pStyle[@w:val="LayoutReferencesStart"]')),None)
  if bib is not None:
   add_boundary_before(bib,current);current=clean_section(sections[2],kind='continuous',columns=1)
  body.replace(body.find(qn('w:sectPr')),current)
 elif front=='paper-cover':
  body.replace(body.find(qn('w:sectPr')),clean_section(sections[1],columns=1,restart=1))
 # Explicit semantic wide region -> landscape section, independent of title text.
 for p in list(doc.paragraphs):
  kind=p._p.xpath('./w:pPr/w:pStyle/@w:val')
  if kind and kind[0]=='LayoutReferencesStart':
   p._p.getparent().remove(p._p);continue
  if kind and kind[0] in ['LayoutLandscapeStart','LayoutLandscapeEnd']:
   if cfg['landscape_regions']:
    # Section 2 in the source is landscape; its next-page transition is kept.
    idx=1 if kind[0]=='LayoutLandscapeStart' else 2
    node=clean_section(sections[idx],columns=1)
    if kind[0]=='LayoutLandscapeEnd':
     size=node.find(qn('w:pgSz'))
     if size is not None:
      size.set(qn('w:w'),str(Cm(29.7).twips));size.set(qn('w:h'),str(Cm(21).twips));size.set(qn('w:orient'),'landscape')
    p._p.clear();props=OxmlElement('w:pPr');props.append(node);p._p.append(props)
   else:p._p.getparent().remove(p._p)
 # A section's properties live at its end, so infer each block's actual region.
 current_section={};nodes=list(body);next_sec=body.find(qn('w:sectPr'))
 for node in reversed(nodes):
  sec=node.find('./'+qn('w:pPr')+'/'+qn('w:sectPr')) if node.tag==qn('w:p') else None
  if sec is not None:next_sec=sec
  current_section[node]=next_sec
 def available(node):
  section=current_section.get(node,body.find(qn('w:sectPr')));size=section.find(qn('w:pgSz'));mar=section.find(qn('w:pgMar'));cols=section.find(qn('w:cols'))
  width=int(size.get(qn('w:w')))-int(mar.get(qn('w:left')))-int(mar.get(qn('w:right')))
  count=int(cols.get(qn('w:num'),'1')) if cols is not None else 1
  gap=int(cols.get(qn('w:space'),'720')) if cols is not None else 720
  return (width-gap*(count-1))/count
 for table in doc.tables:
  if cfg['table_properties']:
   props=etree.fromstring(cfg['table_properties'].encode())
   for tag in ['tblBorders','tblCellMar','tblLook','jc']:
    child=props.find(qn('w:'+tag))
    if child is not None:replace_child(table._tbl.tblPr,child)
  width=available(table._tbl);count=len(table.columns);weights=[1.3,1,2.8,1] if count==4 else [1.1,1,2]
  if len(weights)!=count:weights=[1]*count
  widths=[int(width*w/sum(weights)) for w in weights]
  table.autofit=False
  margins=OxmlElement('w:tblCellMar')
  for edge in ['left','right','top','bottom']:
   n=OxmlElement('w:'+edge);n.set(qn('w:w'),'80' if edge in ['left','right'] else '30');n.set(qn('w:type'),'dxa');margins.append(n)
  replace_child(table._tbl.tblPr,margins)
  for grid,v in zip(table._tbl.tblGrid,widths):grid.set(qn('w:w'),str(v))
  for i,row in enumerate(table.rows):
   row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
   if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
   for cell,v in zip(row.cells,widths):
    cell.width=Cm(v/567)
    for p in cell.paragraphs:
     p.style=lookup(doc,'TableHeader' if i==0 else 'TableBody')
     p.paragraph_format.keep_with_next=(len(table.rows)<=12 and i<len(table.rows)-1)
     p.paragraph_format.alignment=0
     p.paragraph_format.space_after=Pt(3)
     for r in p.runs:r._r.rPr is None or r._r.remove(r._r.rPr)
     if i==0:
      for r in p.runs:r.bold=True
 first_h1=True
 for p in doc.paragraphs:
  for shape in p._p.xpath('.//wp:inline'):
   extent=shape.find(qn('wp:extent'));width=int(extent.get('cx'));height=int(extent.get('cy'));target=int(available(p._p)*635*.94)
   extent.set('cx',str(target));extent.set('cy',str(round(height*target/width)))
   for ext in shape.xpath('.//a:xfrm/a:ext'):ext.set('cx',str(target));ext.set('cy',str(round(height*target/width)))
   p.paragraph_format.first_line_indent=Pt(0);p.alignment=1;p.paragraph_format.keep_with_next=True
  if cfg['chapter_new_page'] and p.style.style_id=='Heading1':
   p.paragraph_format.page_break_before=not first_h1;first_h1=False
  # Formula placement uses native OMML and avoids body indentation.
  if p._p.xpath('.//m:oMathPara'):p.paragraph_format.first_line_indent=Pt(0);p.alignment=1
 # Source-style running matter. Nothing new is added to the journal header.
 for i,section in enumerate(doc.sections):
  section.header.is_linked_to_previous=False;section.footer.is_linked_to_previous=False
  for p in section.header.paragraphs:
   p.clear()
   if not cfg['source_running_header']:
    p.style=lookup(doc,'Normal')
    props=p._p.find(qn('w:pPr'))
    if props is not None:
     for n in list(props.findall(qn('w:pBdr'))):props.remove(n)
  for p in section.footer.paragraphs:p.clear()
  numbered=not(front=='paper-cover' and i==0)
  if front=='paper-cover':
   for old in list(section._sectPr.findall(qn('w:pgNumType'))):section._sectPr.remove(old)
   if i==1:
    number=OxmlElement('w:pgNumType');number.set(qn('w:start'),'1');section._sectPr.append(number)
  if cfg['source_running_header'] and numbered:
   p=section.header.paragraphs[0];p.add_run(meta['course']);p.style=lookup(doc,'Caption');p.alignment=2 if front=='paper-cover' else 0
  if cfg['source_running_footer'] and numbered:
   p=section.footer.paragraphs[0];p.alignment=1
   if front=='paper-cover':p.add_run('— ')
   elif front=='experiment-cover':p.add_run(meta['group']+' '+meta['author']+' | ')
   field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
   if front=='paper-cover':p.add_run(' —')
 # Keep useful view defaults; no source revision history or compatibility banner.
 compat=doc.settings.element.find(qn('w:compat'))
 if compat is not None:
  for item in compat:
   if item.get(qn('w:name'))=='compatibilityMode':item.set(qn('w:val'),'15')
 doc.core_properties.title=meta['title'];doc.core_properties.author=meta['author']
 # East Asian Word defaults can add display spacing not present in source text;
 # content whitespace remains untouched (code indentation is preserved).
 doc.save(output_docx)
 return {'sections':len(doc.sections),'columns':[int(s._sectPr.find(qn('w:cols')).get(qn('w:num'),'1')) for s in doc.sections], 'tables':len(doc.tables),'native_equations':len(doc.element.xpath('//m:oMath'))}

if __name__=='__main__':
 import argparse
 parser=argparse.ArgumentParser(description=__doc__)
 parser.add_argument('input',type=Path);parser.add_argument('output',type=Path)
 parser.add_argument('--layout',type=Path,required=True);parser.add_argument('--metadata',type=Path,required=True)
 args=parser.parse_args();print(json.dumps(apply(args.input,args.output,args.layout,args.metadata),ensure_ascii=False))
