"""Add Pandoc styles to a copy of the NEU bachelor template; preserve its body."""
import json
import subprocess
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from build_word_references import ROOT, lookup, replace_child
from build_neu_journal_reference import promote
from build_word_xml_references import build

SOURCE=Path('D:/MY_DATA/Downloads/毕业设计模板-信息学院202603/2-本科毕业论文模板(2026.3版)V5.docx')
WORK=ROOT/'build/word-xml-review/neu-bachelor'

def main():
    WORK.mkdir(parents=True,exist_ok=True)
    source=Document(SOURCE)
    default=WORK/'pandoc-default.docx'
    default.write_bytes(subprocess.check_output(['quarto','pandoc','--print-default-data-file','reference.docx']))
    seed=Document(default)
    defaults=source.styles.element.find(qn('w:docDefaults'))
    if defaults is not None:
        replace_child(seed.styles.element,defaults)
    # Word implicitly applies the canonical Hyperlink style while rendering TOC.
    # Source uses the localized style ID af0, so carry its format to that alias.
    hyperlink=source.styles['Hyperlink'].element.find(qn('w:rPr'))
    if hyperlink is not None:
        replace_child(lookup(seed,'Hyperlink').element,hyperlink)
    # The supplied template uses Subtitle for level 2, Title for level 3.
    # Alias actual examples rather than relying on their original style names.
    roles={'Title':29,'Author':40,'Abstract':94,'English Abstract':103,
        'Normal':110,'Body Text':110,'First Paragraph':110,'Compact':110,
        'Heading 1':108,'Heading 2':109,'Heading 3':116,
        'Caption':149,'Image Caption':149,'Table Caption':162,
        'Bibliography':306,'Abstract Heading':93,'English Abstract Heading':102}
    for name,index in roles.items():
        promote(source,seed,name,source.paragraphs[index])
    for level in [1,2,3]:
        style=lookup(seed,'Heading '+str(level))
        props=style.element.find(qn('w:pPr'))
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(level-1))
        replace_child(props,outline)
        style.paragraph_format.keep_with_next=True
    for name in ['Compact','Figure','Captioned Figure']:
        style=lookup(seed,name)
        style.paragraph_format.first_line_indent=Pt(0)
        style.paragraph_format.left_indent=Pt(0)
        style.paragraph_format.right_indent=Pt(0)
    table=lookup(seed,'Table')
    borders=source.tables[0]._tbl.tblPr.find(qn('w:tblBorders'))
    if borders is not None:
        props=table.element.find(qn('w:tblPr'))
        if props is None:
            props=OxmlElement('w:tblPr');table.element.append(props)
        replace_child(props,borders)
    try:
        code=lookup(seed,'Source Code')
    except StopIteration:
        code=seed.styles.add_style('Source Code',WD_STYLE_TYPE.PARAGRAPH)
    code.font.name='Consolas';code.font.size=Pt(10)
    code.paragraph_format.first_line_indent=Pt(0)
    seed_path=WORK/'style-seed.docx';seed.save(seed_path)
    specimen=WORK/'specimen.md'
    specimen.write_text('''# 绪论

这是本科毕业论文参考文档的兼容性测试，验证正文与三级标题、列表、公式、表格和代码样式。

## 研究背景

本文保留模板原有的封面、中英文摘要、目录、章节、参考文献和致谢，测试内容另行输出。

### 研究方法

1. 准备数据。
2. 比较分析结果。

$$ L=\\frac{1}{N}\\sum_{i=1}^{N}(y_i-\\hat y_i)^2 $$

```python
def mean(values):
    return sum(values) / len(values)
```

+--------------------------+--------------------------+
| 项目                     | 数值                     |
+==========================+==========================+
| 测试A                    | 0.125                    |
+--------------------------+--------------------------+
| 测试B                    | 0.250                    |
+--------------------------+--------------------------+

: 测试表格

参考链接：[Quarto](https://quarto.org)。
''',encoding='utf-8')
    result=build('neu-bachelor',SOURCE,seed=seed_path,specimen=specimen,preserve_builtins_by_name=True)
    result['sections']=len(source.sections)
    result['visual_review']='pending'
    (WORK/'validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2))

if __name__=='__main__':
    main()
