"""Extract four source-derived reference DOCX files and render one shared specimen.

Run with a Python environment providing python-docx and Pillow. Quarto supplies
Pandoc. Source documents are read only; no LLM is used by this build.
"""
from pathlib import Path
from copy import deepcopy
import hashlib,json,subprocess,sys,shutil
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
from PIL import Image,ImageDraw

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'styles/word'
SOURCES={
 'python-experiment':Path('D:/MY_WORK/Works_Workshop/AfterStory/PythonDL/reports/实验三&实验四 自动化2303 谢怡诺 补充实验完整版.docx'),
 'cad-report':Path('D:/MY_WORK/Works_Workshop/AfterStory/Cad/shou/New folder/第5组_正式提交/案例-5-谢怡诺胡小乐许晋郗杨浩睿.docx'),
 'course-paper':Path('D:/MY_WORK/Works_Workshop/AfterStory/CadPpt/论文/7组自动化2303论文.docx'),
 'electronic-senses':OUT/'electronic-senses/source-converted.docx',
}
LABELS={'python-experiment':'深度学习实验报告','cad-report':'控制实验报告','course-paper':'课程结课论文','electronic-senses':'电子感官论文'}

def lookup(doc,name):
 normalize=lambda value:value.replace(' ','').casefold()
 return next(s for s in doc.styles if normalize(s.name)==normalize(name) or normalize(s.style_id)==normalize(name))

def replace_child(parent,new):
 old=parent.find(new.tag)
 if old is not None:parent.remove(old)
 parent.append(deepcopy(new))

def flatten_style(source,style,kind):
 chain=[]
 while style is not None:
  chain.append(style);style=style.base_style
 result=OxmlElement('w:'+kind)
 defaults=source.styles.element.find(qn('w:docDefaults'))
 if defaults is not None:
  default=defaults.find(qn('w:'+kind+'Default'))
  default=default.find(qn('w:'+kind)) if default is not None else None
  if default is not None:
   for child in default:replace_child(result,child)
 for item in reversed(chain):
  node=item.element.find(qn('w:'+kind))
  if node is not None:
   for child in node:replace_child(result,child)
 return result

def role(doc,target,paragraph=None,source_style=None,omit_layout=False):
 style=lookup(doc,target)
 origin=paragraph.style if paragraph is not None else lookup(doc,source_style or target)
 for kind in ['pPr','rPr']:
  props=flatten_style(doc,origin,kind)
  direct=paragraph._p.find(qn('w:pPr')) if paragraph is not None and kind=='pPr' else None
  if paragraph is not None and kind=='rPr':
   run=max((r for r in paragraph.runs if r.text.strip()),key=lambda r:len(r.text),default=None)
   direct=run._r.find(qn('w:rPr')) if run is not None else None
  if direct is not None:
   for child in direct:
    if child.tag not in [qn('w:pStyle'),qn('w:sectPr')]:replace_child(props,child)
  for bad in ['sectPr','numPr','pStyle','pageBreakBefore']:
   for child in list(props):
    if child.tag==qn('w:'+bad):props.remove(child)
  if omit_layout:
   for child in list(props):
    if child.tag in [qn('w:ind'),qn('w:tabs')]:props.remove(child)
  if kind=='pPr':
   for child in list(props):
    if child.tag==qn('w:rPr'):props.remove(child)
  old=style.element.find(qn('w:'+kind))
  if old is not None:style.element.remove(old)
  style.element.append(props)
 return style

def make_seed(key,path,defaults):
 doc=Document(path);paras=[p for p in doc.paragraphs if p.text.strip()]
 first_section=deepcopy(doc.sections[0]._sectPr)
 # Retain source definitions, add Pandoc's known styles when absent.
 names={s.name.casefold() for s in doc.styles}
 for s in defaults.styles:
  if s.name.casefold() not in names:doc.styles.element.append(deepcopy(s.element));names.add(s.name.casefold())
 # Pandoc-produced sources sometimes use title-case XML names for built-ins;
 # python-docx looks up their canonical lower-case internal names.
 for s in list(doc.styles):s.name=s.name
 if 'Source Code' not in doc.styles:
  code=doc.styles.add_style('Source Code',WD_STYLE_TYPE.PARAGRAPH);code.base_style=doc.styles['Normal'];code.font.name='Consolas';code.font.size=Pt(9)
 if key=='electronic-senses':
  role(doc,'Title',paras[0],omit_layout=True)
  role(doc,'Author',paras[1],omit_layout=True)
  role(doc,'Abstract',next(p for p in paras if p.text.startswith('摘')))
  # The source's body uses Normal; its Body Text style is an unused 7pt style.
  body=next(p for p in paras if p.text.startswith('人根据'))
  for name in ['Normal','Body Text','First Paragraph','Compact']:role(doc,name,body)
  h1=next(p for p in paras if p.text.startswith('1  '))
  h2=next(p for p in paras if p.text.startswith('2.1'))
  role(doc,'Heading 1',h1,omit_layout=True);role(doc,'Heading 2',h2,omit_layout=True)
  caption=next(p for p in paras if p.text.startswith('图1'))
  for name in ['Caption','Image Caption','Table Caption']:role(doc,name,caption,omit_layout=True)
 elif key=='cad-report':
  for name in ['Body Text','First Paragraph','Compact']:role(doc,name,source_style='Normal')
  role(doc,'Title',paras[1],omit_layout=True)
  for name in ['Caption','Image Caption','Table Caption']:role(doc,name,source_style='ReportCaption')
  role(doc,'Source Code',source_style='ReportCode')
 elif key=='course-paper':
  title=next(p for p in paras if p.text.startswith('基于全身'))
  role(doc,'Title',title,omit_layout=True)
  role(doc,'Caption',source_style='Image Caption')
 else:
  for name in ['First Paragraph','Compact']:role(doc,name,source_style='Body Text')
  for name in ['Image Caption','Table Caption']:role(doc,name,source_style='Caption')
 # Promote direct formatting for all visible source roles, not just headings.
 source=Document(path)
 ps=[p for p in source.paragraphs if p.text.strip()]
 def custom(name,p):
  if name not in doc.styles:doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
  return role(doc,name,p)
 if key=='python-experiment':
  for name,index in [('CoverCourse',0),('CoverKind',1),('CoverTopic',2),('CoverGroup',3),('CoverAuthor',5),('CoverDate',7),('CoverNote',8)]:custom(name,ps[index])
 elif key=='cad-report':
  for name,index in [('CoverCourse',0),('CoverKind',1),('CoverTopic',2),('CoverGroup',3),('CoverDate',4)]:custom(name,ps[index])
 elif key=='course-paper':
  for name,index in [('CoverInstitution',0),('CoverCourse',1),('CoverKind',2),('CoverTopic',3),('CoverGroup',4),('CoverSupervisor',5),('CoverDate',6)]:custom(name,ps[index])
 else:
  custom('EnglishTitle',next(p for p in ps if p.text.startswith('An Exploration')))
  custom('EnglishAuthor',next(p for p in ps if p.text.startswith('Xie Yinuo')))
  custom('EnglishAbstract',next(p for p in ps if p.text.startswith('Abstract')))
  # The original has no generic subtitle/date or running header.
  lookup(doc,'Abstract').paragraph_format.space_before=Pt(0)
  h3=lookup(doc,'Heading 3');h3.font.size=Pt(10.5);h3.paragraph_format.line_spacing=1;h3.paragraph_format.space_before=Pt(4);h3.paragraph_format.space_after=Pt(2)
 if key in ['cad-report','course-paper']:
  toc_source=next(p for p in ps if p.text.replace(' ','').replace('　','')=='目录')
  role(doc,'TOC Heading',toc_source)
 # Actual source paragraph overrides are important for compact tables/captions.
 table=source.tables[{'python-experiment':0,'cad-report':1,'course-paper':3}.get(key,0)] if source.tables else None
 if table is not None:
  p=next((p for r in table.rows[1:] for c in r.cells for p in c.paragraphs if p.text),table.cell(0,0).paragraphs[0])
  custom('TableBody',p);custom('TableHeader',table.cell(0,0).paragraphs[0])
 else:
  custom('TableBody',ps[-1]);custom('TableHeader',ps[-1])
  for name in ['TableBody','TableHeader']:lookup(doc,name).font.size=Pt(9)
 for name in ['TableBody','TableHeader']:
  lookup(doc,name).paragraph_format.first_line_indent=Pt(0)
  lookup(doc,name).paragraph_format.keep_with_next=False
  lookup(doc,name).paragraph_format.keep_together=True
 if key=='cad-report':
  custom('Source Code',next(p for p in ps if p.style.name=='ReportCode'))
 code=lookup(doc,'Source Code');code.paragraph_format.alignment=0;code.paragraph_format.first_line_indent=Pt(0);code.paragraph_format.line_spacing=1.1
 code.paragraph_format.space_before=Pt(3);code.paragraph_format.space_after=Pt(3)
 code.font.name='Consolas'
 if key!='cad-report':code.font.size=Pt(9)
 char=lookup(doc,'Verbatim Char');char.font.name='Consolas';char.font.size=Pt(9)
 code.paragraph_format.tab_stops.clear_all()
 if key=='python-experiment':
  cap=next(p for p in ps if p.style.name=='Caption')
  for name in ['Caption','Image Caption','Table Caption']:role(doc,name,cap)
 # Figure/code paragraphs must not inherit body first-line indentation.
 for name in ['Source Code','Figure','Captioned Figure']:
  lookup(doc,name).paragraph_format.first_line_indent=Pt(0)
 for name in ['Heading 1','Heading 2','Heading 3','Table Caption']:
  lookup(doc,name).paragraph_format.keep_with_next=True
 # Promote real table formatting into Pandoc's Table style.
 ts=lookup(doc,'Table');ts.element.set(qn('w:default'),'1')
 if key=='python-experiment':
  actual=doc.tables[0];base=actual.style
  for child in base.element:
   if child.tag not in [qn('w:name'),qn('w:basedOn'),qn('w:next'),qn('w:link')]:replace_child(ts.element,child)
  # Table styles can have multiple conditional definitions.
  for node in list(ts.element.findall(qn('w:tblStylePr'))):ts.element.remove(node)
  for node in base.element.findall(qn('w:tblStylePr')):ts.element.append(deepcopy(node))
 else:
  props=OxmlElement('w:tblPr');borders=OxmlElement('w:tblBorders')
  for edge in ['top','bottom','left','right','insideH','insideV']:
   node=OxmlElement('w:'+edge);node.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');node.set(qn('w:sz'),'6');borders.append(node)
  props.append(borders);replace_child(ts.element,props)
  conditional=OxmlElement('w:tblStylePr');conditional.set(qn('w:type'),'firstRow')
  cp=OxmlElement('w:tcPr');cb=OxmlElement('w:tcBorders');edge=OxmlElement('w:bottom');edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');cb.append(edge);cp.append(cb);conditional.append(cp)
  rp=OxmlElement('w:rPr');rp.append(OxmlElement('w:b'));conditional.append(rp)
  for node in list(ts.element.findall(qn('w:tblStylePr'))):ts.element.remove(node)
  ts.element.append(conditional)
 # Transfer style definitions to a clean package: old media, annotations,
 # external links and source document relationships are not reusable styling.
 source_doc=doc
 doc=Document(OUT/'pandoc-default.docx')
 doc.part._styles_part._element=deepcopy(source_doc.styles.element)
 # Preserve typography-relevant compatibility and grid behavior as well.
 for tag in ['compat','themeFontLang','characterSpacingControl','doNotUseHTMLParagraphAutoSpacing','defaultTabStop']:
  node=source.settings.element.find(qn('w:'+tag))
  if node is not None:replace_child(doc.settings.element,node)
 for reltype in ['http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme','http://schemas.openxmlformats.org/officeDocument/2006/relationships/fontTable']:
  src=next((r.target_part for r in source_doc.part.rels.values() if r.reltype==reltype),None)
  dst=next((r.target_part for r in doc.part.rels.values() if r.reltype==reltype),None)
  if src is not None and dst is not None:
   if hasattr(dst,'_element'):
    from lxml import etree
    dst._element=etree.fromstring(src.blob)
   else:dst._blob=src.blob
 for child in list(first_section):
  if child.tag in [qn('w:headerReference'),qn('w:footerReference')]:first_section.remove(child)
 # Source identity and existing body are not part of the reusable reference.
 for child in list(doc.element.body):doc.element.body.remove(child)
 doc.element.body.append(first_section)
 sec=doc.sections[0]
 sec.different_first_page_header_footer=False
 # A generic PAGE field preserves the footer's role without old identities.
 for container in [sec.header,sec.footer,sec.first_page_header,sec.first_page_footer,sec.even_page_header,sec.even_page_footer]:
  for p in container.paragraphs:
   for child in list(p._p):
    if child.tag!=qn('w:pPr'):p._p.remove(child)
  for table in list(container.tables):table._tbl.getparent().remove(table._tbl)
 if key!='electronic-senses':sec.header.paragraphs[0].add_run(LABELS[key])
 p=sec.footer.paragraphs[0];p.alignment=1
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');p._p.append(field)
 for node in list(sec._sectPr):
  if node.tag in [qn('w:pgNumType'),qn('w:titlePg')]:sec._sectPr.remove(node)
 doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title=LABELS[key]+'样式参考'
 dest=OUT/key;dest.mkdir(parents=True,exist_ok=True);seed=dest/'style-seed.docx';doc.save(seed)
 return seed

def main():
 from apply_word_layout import extract_layout,apply
 sys.stdout.reconfigure(encoding='utf8');OUT.mkdir(exist_ok=True)
 default=OUT/'pandoc-default.docx'
 with default.open('wb') as stream:
  subprocess.run(['quarto','pandoc','--print-default-data-file','reference.docx'],stdout=stream,check=True)
 defaults=Document(default)
 # Identical original schematic for all four specimens (no experimental result).
 im=Image.new('RGB',(1000,220),'white');dr=ImageDraw.Draw(im)
 for i,label in enumerate(['INPUT','PROCESS','OUTPUT']):
  x=30+i*340;dr.rounded_rectangle((x,55,x+260,160),radius=14,fill='#eaf1f6',outline='#31566f',width=3);dr.text((x+90,100),label,fill='#183a51')
  if i<2:dr.line((x+265,105,x+325,105),fill='#25847d',width=5);dr.polygon([(x+325,105),(x+310,95),(x+310,115)],fill='#25847d')
 im.crop((0,35,1000,180)).save(OUT/'sample-flow.png')
 records=[]
 for key,path in SOURCES.items():
  seed=make_seed(key,path,defaults);dest=seed.parent
  layout=dest/'layout.json';extract_layout(key,path,layout)
  work=ROOT/'build/word-reference-review'/key;work.mkdir(parents=True,exist_ok=True)
  raw=work/'raw.docx';reviewed=dest/'reference-reviewed.docx'
  subprocess.run(['quarto','pandoc',str(OUT/'specimen.md'),'--metadata-file='+str(OUT/'specimen-metadata.json'),'--from=markdown+tex_math_dollars','--standalone','--number-sections','--lua-filter='+str(ROOT/'scripts/word-layout-markers.lua'),'--resource-path='+str(OUT),'--reference-doc='+str(seed),'--output='+str(raw)],check=True)
  structure=apply(raw,reviewed,layout,OUT/'specimen-metadata.json')
  updated=True
  try:shutil.copyfile(reviewed,dest/'reference.docx')
  except PermissionError:updated=False
  final=Document(reviewed)
  assert len(final.tables)>=1 and len(final.inline_shapes)==2
  assert final.element.xpath('//m:oMath'),'Native equation missing'
  record=dict(name=key,label=LABELS[key],source=str(path),source_sha256=hashlib.sha256(path.read_bytes()).hexdigest(),reference=str(reviewed.relative_to(ROOT)),original_reference_updated=updated,paragraphs=len(final.paragraphs),tables=len(final.tables),equations=len(final.element.xpath('//m:oMath')),structure=structure)
  records.append(record);print(json.dumps(record,ensure_ascii=False))
 (OUT/'manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')

if __name__=='__main__':main()
