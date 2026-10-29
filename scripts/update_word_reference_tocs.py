"""Materialize a native Word TOC's cached result from actual PDF page positions."""
from pathlib import Path
import re,json,sys
sys.stdout.reconfigure(encoding='utf8')
import pymupdf as fitz
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT,WD_TAB_LEADER
from docx.shared import Cm,Pt
from build_word_references import OUT,ROOT,lookup

def update(key):
 path=OUT/key/'reference-reviewed.docx';pdf=ROOT/'build/word-reference-review'/f'{key}.pdf'
 d=Document(path);pages=fitz.open(pdf)
 normalize=lambda t:re.sub(r'\s+','',t)
 anchors=[p for p in d.paragraphs if p._p.xpath('.//w:fldSimple[starts-with(@w:instr,"TOC")] | .//w:instrText[starts-with(text(),"TOC")]')]
 if not anchors:return {'name':key,'toc':'not_applicable'}
 anchor=anchors[0]._p
 old=[p._p for p in d.paragraphs if p._p.xpath('.//w:bookmarkStart[starts-with(@w:name,"_layout_toc_")]')]
 headings=[p for p in d.paragraphs if p.style.style_id in ['Heading1','Heading2']]
 entries=[]
 for p in headings:
  needle=normalize(p.text);found=None
  for i,page in enumerate(pages):
   if i<2:continue  # cover and TOC are not body matches
   if needle in normalize(page.get_text()):found=i;break
  if found is None:raise ValueError(f'PDF heading not found: {key}: {p.text}')
  logical=found-1 if key=='course-paper' else found+1
  entries.append((re.sub(r'\s+',' ',p.text).strip(),logical,1 if p.style.style_id=='Heading1' else 2))
 # Insert new cached paragraphs before removing old ones; the first old entry
 # may itself be our anchor on subsequent passes.
 added=[]
 for i,(text,page,level) in enumerate(entries):
  name=f'Review TOC {level}'
  try:style=lookup(d,name)
  except StopIteration:
   style=d.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH);style.base_style=lookup(d,'Normal');style.font.size=Pt(11)
  p=d.add_paragraph(style=style);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.left_indent=Cm(.5 if level==2 else 0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(3);p.paragraph_format.alignment=0;p.paragraph_format.keep_with_next=False
  p.paragraph_format.tab_stops.clear_all()
  width=(d.sections[0].page_width-d.sections[0].left_margin-d.sections[0].right_margin)/360000
  p.paragraph_format.tab_stops.add_tab_stop(Cm(width),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
  mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),str(10000+i));mark.set(qn('w:name'),f'_layout_toc_{i}');p._p.append(mark)
  p.add_run(text+'\t'+str(page))
  end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(10000+i));p._p.append(end)
  anchor.addprevious(p._p);added.append(p)
 # True TOC field with cached results, not an unrelated static list.
 first=added[0]._p
 r=OxmlElement('w:r');begin=OxmlElement('w:fldChar');begin.set(qn('w:fldCharType'),'begin');r.append(begin)
 instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text='TOC \\o "1-2" \\h \\z \\u';r.append(instr)
 sep=OxmlElement('w:fldChar');sep.set(qn('w:fldCharType'),'separate');r.append(sep);first.insert(1,r)
 r=OxmlElement('w:r');end=OxmlElement('w:fldChar');end.set(qn('w:fldCharType'),'end');r.append(end);added[-1]._p.append(r)
 for node in dict.fromkeys([anchor,*old]):
  if node.getparent() is not None:node.getparent().remove(node)
 d.save(path)
 return {'name':key,'toc':'updated','entries':entries}

if __name__=='__main__':
 result=[update(k) for k in ['cad-report','course-paper']]
 (OUT/'toc-validation.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps(result,ensure_ascii=False))
