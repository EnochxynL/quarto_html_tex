"""Check content retention, sections, PDF bounds, source hashes and TOC pages."""
from pathlib import Path
import json,hashlib,re,sys,shutil
import pymupdf as fitz
from docx import Document
from docx.oxml.ns import qn
from build_word_references import ROOT,OUT
sys.stdout.reconfigure(encoding='utf8')

def main(visual_review_confirmed=False):
 records=json.loads((OUT/'manifest.json').read_text(encoding='utf8'));checks=[]
 tocs={v['name']:v for v in json.loads((OUT/'toc-validation.json').read_text(encoding='utf8'))}
 for entry in records:
  key=entry['name'];folder=OUT/key;path=folder/'reference-reviewed.docx';d=Document(path)
  raw=Document(ROOT/'build/word-reference-review'/key/'raw.docx')
  start=next(i for i,p in enumerate(raw.paragraphs) if p._p.xpath('./w:pPr/w:pStyle[@w:val="LayoutBodyStart"]'))
  expected=[p.text for p in raw.paragraphs[start+1:] if p.text]
  actual=[p.text for p in d.paragraphs if p.text];actual=actual[actual.index(expected[0]):]
  assert expected==actual,(key,'body text changed')
  assert len(d.element.xpath('//m:oMath'))==2
  assert len(d.inline_shapes)==2
  columns=[int(s._sectPr.find(qn('w:cols')).get(qn('w:num'),'1')) for s in d.sections]
  assert columns==([1,2,1] if key=='electronic-senses' else [1]*len(d.sections))
  assert sum(s.page_width>s.page_height for s in d.sections)==(1 if key=='course-paper' else 0)
  assert hashlib.sha256(Path(entry['source']).read_bytes()).hexdigest()==entry['source_sha256']
  pdf=fitz.open(ROOT/'build/word-reference-review'/f'{key}.pdf')
  for page in pdf:
   assert len(page.get_text().strip())>20,(key,'empty page')
   assert '\ufffd' not in page.get_text(),(key,'replacement glyph')
   for b in page.get_text('blocks'):
    assert b[0]>=-1 and b[1]>=-1 and b[2]<=page.rect.width+1 and b[3]<=page.rect.height+1,(key,'outside page',b[:4])
  if key in tocs:
   normalize=lambda t:re.sub(r'\s+','',t)
   toc_text=normalize(pdf[1].get_text())
   for title,logical,level in tocs[key]['entries']:
    physical=logical+1 if key=='course-paper' else logical-1
    assert normalize(title) in normalize(pdf[physical].get_text()),(key,'wrong TOC page',title)
    assert normalize(title) in toc_text,(key,'missing TOC entry',title)
  # Publish a matched DOCX/PDF pair next to each style candidate.
  shutil.copyfile(ROOT/'build/word-reference-review'/f'{key}.pdf',folder/'reference-reviewed.pdf')
  try:shutil.copyfile(path,folder/'reference.docx');entry['original_reference_updated']=True
  except PermissionError:entry['original_reference_updated']=False
  checks.append(dict(name=key,status='passed',pdf_pages=len(pdf),sections=len(d.sections),columns=columns,source_unchanged=True,body_retained=True,native_equations=2,visual_review='confirmed_after_page_review' if visual_review_confirmed else 'pending',docx_sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
 (OUT/'manifest.json').write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf8')
 (OUT/'validation.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf8')
 print(json.dumps(checks,ensure_ascii=False,indent=2))
if __name__=='__main__':
 import argparse
 parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--visual-review-confirmed',action='store_true',help='Only set after actually inspecting the current rendered pages.')
 main(parser.parse_args().visual_review_confirmed)
